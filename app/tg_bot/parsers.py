"""Парсинг файлов резюме: PDF и DOCX."""
from __future__ import annotations

import io
from PyPDF2 import PdfReader
from docx import Document


def extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    chunks: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            chunks.append(text)
    return "\n".join(chunks).strip()


def extract_text_from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts).strip()


def parse_document(file_name: str, data: bytes) -> str:
    name = (file_name or "").lower()
    if name.endswith(".pdf"):
        return extract_text_from_pdf(data)
    if name.endswith(".docx"):
        return extract_text_from_docx(data)
    raise ValueError("unsupported_format")
